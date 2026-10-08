#!/usr/bin/env python3
"""OWNER_DECISION_L2_018: mechanical pos mapping for the L2_017 candidates: حرف جر→حرف ، اسم مبني→مبني ، صفة جامدة→اسم جامد.
«صفة» and «حال» are NOT mapped (flag POS_NOT_CLOSED stays; owner decides in review). Root/wazn cleared when the new pos is حرف/مبني.
Then writes 03_lus_ar/l2_017_review.xlsx with ALL candidates (flagged first) for owner review."""
import csv, hashlib, json, pathlib, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
ROOT = pathlib.Path(__file__).resolve().parents[1]
P = ROOT / "03_lus_ar/l2_017_cand_norm.csv"
assert hashlib.sha256(P.read_bytes()).hexdigest() == "92813e3c837f8948861f985afe823e9972404d22a4fbb8bb7794d6ed25e9de0d", "STOP norm sha"
MAP = {"حرف جر": "حرف", "اسم مبني": "مبني", "صفة جامدة": "اسم جامد"}
POS = [l.strip() for l in open(ROOT / "03_lus_ar/pos_ar_closed.txt", encoding="utf-8") if l.strip()]
rows = list(csv.DictReader(open(P, encoding="utf-8"))); cols = list(rows[0]); st = collections.Counter()
for r in rows:
    if r["pos_ar"] in MAP:
        st[f"mapped {r['pos_ar']}→{MAP[r['pos_ar']]}"] += 1; r["pos_ar"] = MAP[r["pos_ar"]]
        fl = set(filter(None, r["flags"].split(";"))) - {"POS_NOT_CLOSED"}
        if r["pos_ar"] in ("حرف", "مبني") and (r["root"] or r["wazn"] or r["wazn_note"]):
            r["root"] = r["wazn"] = r["wazn_note"] = ""; fl -= {"ROOT_ALIF","ROOT_LEN","ROOT_MISSING","ROOT_NORMALIZED","WAZN_NOT_IN_INV","WAZN_MISSING","WAZN_LEMMA_MISMATCH"}; fl.add("ROOT_WAZN_CLEARED_NO_ROOT_POS")
        r["flags"] = ";".join(sorted(fl)); r["note"] += " | POS ← L2_018"
out = ROOT / "03_lus_ar/l2_017_cand_norm2.csv"
with open(out, "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(rows)
# review workbook
fr = {}
for l in open(ROOT / "release/layer1_v3/frames_ar.jsonl", encoding="utf-8"):
    d = json.loads(l); fr[str(d["frame_id"])] = (d["name"], d.get("name_ar", ""))
en = {r["lu_id"]: f"{r['lemma']}.{r['pos']}" for r in csv.DictReader(open(ROOT / "03_lus_ar/lus_template.csv", encoding="utf-8-sig"))}
H = ["ar_lu_id","frame","frame_ar","lemma_ar","pos_ar","root","wazn","wazn_note","definition_ar","en_lemmas","en_lu_ids","flags","status","note","owner_choice"]
srt = sorted(rows, key=lambda r: (r["flags"] == "", int(r["frame_id"]), r["ar_lu_id"]))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "l2_017"; ws.sheet_view.rightToLeft = True
ws.append(H)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
yel = PatternFill("solid", fgColor="FFF2CC")
for r in srt:
    n, na = fr.get(r["frame_id"], ("?", ""))
    ws.append([r["ar_lu_id"], n, na, r["lemma_ar"], r["pos_ar"], r["root"], r["wazn"], r["wazn_note"], r["definition_ar"],
               ", ".join(en.get(i, i) for i in r["en_lu_ids"].split(";")), r["en_lu_ids"], r["flags"], r["status"], r["note"], ""])
    if r["flags"]:
        ws.cell(ws.max_row, 12).fill = yel
for col, wdt in zip("ABCDEFGHIJKLMNO", [10,24,20,16,11,7,12,12,45,30,14,30,11,24,14]): ws.column_dimensions[col].width = wdt
ws.freeze_panes = "E2"; ws.auto_filter.ref = ws.dimensions
g = wb.create_sheet("flags_legend")
for k, v in [("AL_STRIPPED","أُزيلت أل التعريف آليًا — تحقّق"),("ROOT_NORMALIZED","وُحّد الجذر (همزات/ى/علامات)"),("ROOT_ALIF","ألف وسطية/طرفية في الجذر — أصلها مجهول"),
             ("ROOT_LEN","طول الجذر ليس 3 أو 4"),("ROOT_MISSING","لا جذر"),("WAZN_NOT_IN_INV","الوزن ليس في المخزون المعتمد → نُقل إلى wazn_note"),
             ("WAZN_MISSING","لا وزن"),("WAZN_LEMMA_MISMATCH","الوزن مع الجذر لا يطابق هيكل الكلمة (جذور صحيحة فقط)"),("POS_NOT_CLOSED","القسم خارج القائمة المغلقة (صفة/حال) — قرّر"),
             ("MUBALAGHA_NOT_FAAAL","صيغة مبالغة على غير فَعَّال"),("MARRA_NOT_FALA","اسم مرة على غير فَعْلَة"),("UNVOCALIZED","الكلمة غير مشكولة"),
             ("ROOT_WAZN_CLEARED_NO_ROOT_POS","حرف/مبني/تركيب: حُذف الجذر والوزن"),("AL_UNRESOLVED","أل لم تُزل آليًا"),
             ("owner_choice","اكتب: = (اعتماد كما هو) أو الصيغة المصحّحة أو X (رفض)")]: g.append([k, v])
g.column_dimensions["A"].width = 32; g.column_dimensions["B"].width = 70; g.sheet_view.rightToLeft = True
x = ROOT / "03_lus_ar/l2_017_review.xlsx"; wb.save(x)
for k, v in sorted(st.items()): print(k, "=", v)
fc = collections.Counter(f for r in rows for f in filter(None, r["flags"].split(";")))
print("ar_lus =", len(rows), "| unflagged =", sum(1 for r in rows if not r["flags"]))
for k, v in sorted(fc.items()): print("flag_" + k, "=", v)
print("pos_not_closed_left =", dict(collections.Counter(r["pos_ar"] for r in rows if r["pos_ar"] not in POS)))
print("sha256 l2_017_cand_norm2.csv =", hashlib.sha256(out.read_bytes()).hexdigest())
print("review rows =", ws.max_row - 1)
